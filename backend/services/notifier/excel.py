"""Excel-отчёт с диаграммами и рекомендациями (openpyxl, нативные диаграммы Excel)."""
import io
from datetime import timedelta

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import analytics as an, fmt
from .reports import LVL, STATUS_ICON, month_goals
from .service import Bundle

HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
HEAD_FONT = Font(bold=True, color="FFFFFF")
MONEY = '# ##0 "₽"'
LEVEL_FILL = {"warn": "FDE68A", "info": "DBEAFE", "ok": "DCFCE7"}


def _r(kop: int | None) -> float | None:
    return None if kop is None else round(kop / 100, 2)


def _head(ws, row: int, titles: list[str]):
    for i, t in enumerate(titles, 1):
        c = ws.cell(row=row, column=i, value=t)
        c.fill, c.font = HEAD_FILL, HEAD_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _widths(ws, widths: list[int]):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def build_xlsx(b: Bundle, kind: str = "week") -> bytes:
    wb = Workbook()
    _summary(wb.active, b, kind)
    _forecast(wb.create_sheet("Прогноз"), b)
    _flows(wb.create_sheet("Поступления и расходы"), b)
    _plan_fact(wb.create_sheet("План-факт"), b)
    _goals(wb.create_sheet("Цели"), b)
    _tips(wb.create_sheet("Рекомендации"), b)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _summary(ws, b: Bundle, kind: str):
    ws.title = "Сводка"
    ws["A1"] = f"Финансовый отчёт · {b.company}"
    ws["A1"].font = Font(bold=True, size=16)
    ws["A2"] = f"Сформирован {fmt.day(b.today)} {b.today.year}"
    ws["A2"].font = Font(italic=True, color="6B7280")
    w, m = b.week, b.month_last
    rows = [
        ("Остаток по данным", _r(b.fc.opening), f"на {fmt.day(b.fc.as_of)}"),
        ("Уровень риска", LVL[b.risk.level][2:].strip(), ""),
        ("Вероятность кассового разрыва", b.fc.prob, "за горизонт прогноза"),
        ("Дата возможного разрыва", fmt.day(b.fc.break_date) if b.fc.break_date else "не ожидается", ""),
        (f"Поступления за {fmt.period(w['start'], w['end'])}", _r(w["income"]), "прошлая неделя"),
        (f"Расходы за {fmt.period(w['start'], w['end'])}", _r(w["expense"]), "прошлая неделя"),
        (f"Поступления за {fmt.month_title(m['start'])}", _r(m["income"]), "прошлый месяц"),
        (f"Расходы за {fmt.month_title(m['start'])}", _r(m["expense"]), "прошлый месяц"),
        ("Чистый поток за месяц", _r(m["net"]), ""),
    ]
    _head(ws, 4, ["Показатель", "Значение", "Пояснение"])
    for i, (a, v, c) in enumerate(rows, 5):
        ws.cell(row=i, column=1, value=a)
        cell = ws.cell(row=i, column=2, value=v)
        if isinstance(v, float) and v <= 1 and "Вероятность" in a:
            cell.number_format = "0%"
        elif isinstance(v, float):
            cell.number_format = MONEY
        ws.cell(row=i, column=3, value=c)
    r = 5 + len(rows) + 1
    ws.cell(row=r, column=1, value="План на месяц").font = Font(bold=True, size=12)
    for i, g in enumerate(month_goals(b), r + 1):
        ws.cell(row=i, column=1, value=f"{i - r}. {g}")
    _widths(ws, [46, 22, 34])


def _forecast(ws, b: Bundle):
    _head(ws, 1, ["Дата", "Худший (p10)", "Медиана (p50)", "Лучший (p90)", "Вероятность минуса"])
    for i, d in enumerate(b.fc.days, 2):
        ws.cell(row=i, column=1, value=d.date).number_format = "DD.MM.YYYY"
        for j, v in enumerate((d.p10, d.p50, d.p90), 2):
            ws.cell(row=i, column=j, value=_r(v)).number_format = MONEY
        ws.cell(row=i, column=5, value=d.prob_neg).number_format = "0%"
    n = len(b.fc.days)
    if n:
        ch = LineChart()
        ch.title, ch.height, ch.width = "Прогноз остатка", 9, 22
        ch.y_axis.title, ch.x_axis.number_format = "₽", "DD.MM"
        ch.add_data(Reference(ws, min_col=2, max_col=4, min_row=1, max_row=n + 1), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=n + 1))
        ws.add_chart(ch, "G2")
    _widths(ws, [14, 16, 16, 16, 18])


def _flows(ws, b: Bundle):
    m = b.month_last
    _head(ws, 1, ["Неделя", "Поступления", "Расходы", "Чистый поток"])
    flows = an.weekly_flows(b.txs, m["start"], m["end"])
    for i, f in enumerate(flows, 2):
        ws.cell(row=i, column=1, value=f"{f['start'].day}–{f['end'].day} {fmt.MONTHS[f['end'].month - 1]}")
        ws.cell(row=i, column=2, value=_r(f["income"])).number_format = MONEY
        ws.cell(row=i, column=3, value=_r(f["expense"])).number_format = MONEY
        ws.cell(row=i, column=4, value=f"=B{i}-C{i}").number_format = MONEY
    if flows:
        ch = BarChart()
        ch.title, ch.height, ch.width = f"Поступления и расходы по неделям · {fmt.month_title(m['start'])}", 8, 16
        ch.add_data(Reference(ws, min_col=2, max_col=3, min_row=1, max_row=len(flows) + 1), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=len(flows) + 1))
        ws.add_chart(ch, "F2")
    base = len(flows) + 4
    _head(ws, base, ["Категория расходов", "Сумма", "Доля"])
    cats = list(m["by_cat_out"].items())
    for i, (c, v) in enumerate(cats, base + 1):
        ws.cell(row=i, column=1, value=fmt.category(c))
        ws.cell(row=i, column=2, value=_r(v)).number_format = MONEY
        ws.cell(row=i, column=3, value=(v / m["expense"]) if m["expense"] else 0).number_format = "0%"
    if cats:
        ch = BarChart()
        ch.type, ch.title, ch.height, ch.width = "bar", "Структура расходов за месяц", 8, 16
        ch.add_data(Reference(ws, min_col=2, min_row=base, max_row=base + len(cats)), titles_from_data=True)
        ch.set_categories(Reference(ws, min_col=1, min_row=base + 1, max_row=base + len(cats)))
        ch.legend = None
        ws.add_chart(ch, f"F{base}")
    _widths(ws, [30, 18, 18, 18])


def _plan_fact(ws, b: Bundle):
    _head(ws, 1, ["Дата", "Описание", "План", "Факт", "Отклонение", "Статус"])
    names = {"paid": "оплачено", "deviation": "с отклонением", "missed": "не состоялось", "pending": "ожидается"}
    if not b.pf:
        ws["A2"] = "Данных план-факта нет"
        return
    items = an.plan_fact_stats(b.pf, b.today - timedelta(days=60), b.today)["items"]
    for i, it in enumerate(sorted(items, key=lambda x: x["date"]), 2):
        ws.cell(row=i, column=1, value=it["date"]).number_format = "DD.MM.YYYY"
        ws.cell(row=i, column=2, value=it["title"])
        ws.cell(row=i, column=3, value=_r(it["plan"])).number_format = MONEY
        ws.cell(row=i, column=4, value=_r(it["fact"])).number_format = MONEY
        if it["fact"] is not None:
            ws.cell(row=i, column=5, value=f"=D{i}-C{i}").number_format = MONEY
        ws.cell(row=i, column=6, value=names.get(it["status"], it["status"]))
    _widths(ws, [14, 40, 16, 16, 16, 18])


def _goals(ws, b: Bundle):
    _head(ws, 1, ["Спринт", "Цель", "Целевая сумма", "Срок", "Статус"])
    rows = [(b.goals_prev_week, g) for g in b.goals_prev] + [(b.today.isoformat(), g) for g in b.goals_now]
    names = {"open": "в работе", "done": "выполнено", "partial": "частично", "failed": "не выполнено", "skipped": "пропущено"}
    for i, (w, g) in enumerate(rows, 2):
        ws.cell(row=i, column=1, value=w)
        ws.cell(row=i, column=2, value=g["title"]).alignment = Alignment(wrap_text=True)
        ws.cell(row=i, column=3, value=_r(g.get("target_kop"))).number_format = MONEY
        ws.cell(row=i, column=4, value=g.get("due"))
        ws.cell(row=i, column=5, value=f"{STATUS_ICON.get(g.get('status', 'open'), '')} {names.get(g.get('status', 'open'), '')}".strip())
    if b.vel is not None:
        ws.cell(row=len(rows) + 3, column=2, value="Выполнение целей прошлого спринта").font = Font(bold=True)
        ws.cell(row=len(rows) + 3, column=3, value=b.vel).number_format = "0%"
    _widths(ws, [14, 70, 18, 14, 18])


def advice_tips(a: dict | None) -> list[dict]:
    if not a:
        return []
    tips = [{"level": "info", "title": "Консультант Hermes: итог", "text": a["summary"]}]
    tips += [{"level": "warn", "title": f"Приоритет {i}", "text": x} for i, x in enumerate(a["priorities"], 1)]
    tips += [{"level": "info", "title": "Действие на неделю", "text": x} for x in a["actions"]]
    if a["consulting"]:
        tips.append({"level": "ok", "title": "Совет по управлению деньгами", "text": a["consulting"]})
    return tips


def _tips(ws, b: Bundle):
    _head(ws, 1, ["Что делать", "Почему и как"])
    for i, t in enumerate(advice_tips(b.advice) + b.tips, 2):
        a, c = ws.cell(row=i, column=1, value=t["title"]), ws.cell(row=i, column=2, value=t["text"])
        a.font = Font(bold=True)
        a.fill = PatternFill("solid", fgColor=LEVEL_FILL.get(t["level"], "FFFFFF"))
        a.alignment = c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[i].height = max(30, 15 * (len(t["text"]) // 90 + 1))
    _widths(ws, [38, 100])
