"""Разбор банковской выписки CSV/XLSX/1C в каноничный формат.

Форматы: CSV/TXT/XLSX с заголовками; текстовый обмен 1С «1CClientBankExchange» (.txt/.1c,
cp1251/cp866/utf-8) — идея и разбор: проект gleb (swimshadow), services/ingest/src/parse.rs.

Принципы: ничего не исправляем молча; сомнительные строки идут в errors.
Суммы — целые копейки (строкой), без float-арифметики.
"""
import csv
import io
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from .categorize import categorize_ex
from .taxonomy import classify
from .tail_model import apply_tail_model
from .tax_regime import guess_tax_regime

MAX_ROWS = 200_000

ALIASES = {
    "date": ["дата операции", "дата", "date", "operation_date", "дата проводки", "дата платежа"],
    "amount": ["сумма", "amount", "сумма операции"],
    "debit": ["списание", "дебет", "расход", "debit", "сумма списания"],
    "credit": ["поступление", "кредит", "приход", "credit", "сумма поступления"],
    "direction": ["направление", "тип операции", "тип", "direction", "type"],
    "counterparty": ["контрагент", "наименование контрагента", "наименование", "counterparty", "название контрагента", "counterparty_name"],
    "inn": ["инн", "инн контрагента", "inn", "counterparty_inn"],
    "purpose": ["назначение платежа", "назначение", "purpose", "description", "описание", "payment_purpose"],
    "source_id": ["operation_id", "id операции", "номер операции"],
    "company": ["company_id", "код компании"],
    "balance": ["остаток", "остаток на конец дня", "остаток на счете", "balance", "balance_after", "остаток на счёте", "balance_eod"],
}
IN_WORDS = {"in", "credit", "приход", "поступление", "кредит", "income", "+", "входящий"}
OUT_WORDS = {"out", "debit", "расход", "списание", "дебет", "expense", "-", "исходящий"}


class RowError(Exception):
    def __init__(self, code: str):
        self.code = code


def _norm(h: str) -> str:
    return re.sub(r"\s+", " ", str(h or "").strip().lower().replace("﻿", ""))


def map_columns(headers: list[str]) -> dict[str, int]:
    normed = [_norm(h) for h in headers]
    out: dict[str, int] = {}
    for key, names in ALIASES.items():
        for n in names:
            if n in normed:
                out[key] = normed.index(n)
                break
    return out


def parse_money(raw) -> int:
    """'1 234,56' | 1234.56 | '-1234.5' -> копейки."""
    if raw is None or str(raw).strip() == "":
        raise RowError("empty_amount")
    s = str(raw).strip().replace(" ", "").replace(" ", "")
    if "," in s and "." in s:
        s = s.replace(",", "")
    else:
        s = s.replace(",", ".")
    try:
        d = Decimal(s)
    except InvalidOperation:
        raise RowError("bad_amount")
    if not d.is_finite():
        raise RowError("bad_amount")
    return int((d * 100).to_integral_value())


def parse_date(raw) -> date:
    if isinstance(raw, datetime):
        return raw.date()
    if isinstance(raw, date):
        return raw
    s = str(raw or "").strip()
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d.%m.%y", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%d %H:%M:%S", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    raise RowError("bad_date")


def read_rows(filename: str, data: bytes) -> list[list]:
    name = filename.lower()
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)  # формулы не исполняются
        ws = wb.active
        rows = []
        for r in ws.iter_rows(values_only=True):
            rows.append(list(r))
            if len(rows) > MAX_ROWS + 1:
                raise ValueError("too_many_rows")
        return rows
    if name.endswith(".csv") or name.endswith(".txt"):
        for enc in ("utf-8-sig", "cp1251"):
            try:
                text = data.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise ValueError("bad_encoding")
        sample = text[:4096]
        delim = max(";,\t|", key=sample.count)
        rows = []
        for r in csv.reader(io.StringIO(text), delimiter=delim):
            rows.append(r)
            if len(rows) > MAX_ROWS + 1:
                raise ValueError("too_many_rows")
        return rows
    raise ValueError("unsupported_type")


def parse_statement(filename: str, data: bytes) -> dict:
    name = (filename or "").lower()
    if name.endswith(".1c") or (name.endswith((".txt", ".csv")) and is_1c(data)):
        return parse_1c(data)
    rows = read_rows(filename, data)
    if not rows:
        raise ValueError("empty_file")
    cols = map_columns([str(c) for c in rows[0]])
    if "date" not in cols:
        raise ValueError("no_date_column")
    has_amount = "amount" in cols or ("debit" in cols or "credit" in cols)
    if not has_amount:
        raise ValueError("no_amount_column")
    if "company" in cols:
        # общий файл датасета (all_operations.csv): выписки разных компаний склеивать нельзя — интерфейс делит его сам
        ids = {str(r[cols["company"]]).strip() for r in rows[1:] if len(r) > cols["company"] and str(r[cols["company"]]).strip()}
        if len(ids) > 1:
            raise ValueError("multiple_companies")

    txs, errors = [], []
    for i, row in enumerate(rows[1:], start=2):
        if not any(str(c).strip() for c in row if c is not None):
            continue

        def get(k):
            j = cols.get(k)
            return row[j] if j is not None and j < len(row) else None

        try:
            d = parse_date(get("date"))
            deb_s = str(get("debit") or "").strip()
            cre_s = str(get("credit") or "").strip()
            if deb_s or cre_s:
                deb = parse_money(deb_s) if deb_s else 0
                cre = parse_money(cre_s) if cre_s else 0
                if deb and cre:
                    raise RowError("both_debit_credit")
                signed = cre if cre else -abs(deb)
            else:
                signed = parse_money(get("amount"))
                dir_raw = str(get("direction") or "").strip().lower()
                if dir_raw:
                    if dir_raw in IN_WORDS:
                        signed = abs(signed)
                    elif dir_raw in OUT_WORDS:
                        signed = -abs(signed)
                    else:
                        raise RowError("bad_direction")
                elif "direction" in cols:
                    raise RowError("bad_direction")
            if signed == 0:
                raise RowError("zero_amount")
            balance = None
            if str(get("balance") or "").strip():
                balance = parse_money(get("balance"))
            direction = "in" if signed > 0 else "out"
            sid = str(get("source_id") or "").strip()[:100]
            txs.append(
                _make_tx(
                    i, d, signed, str(get("counterparty") or ""), str(get("inn") or ""),
                    str(get("purpose") or ""), balance, sid,
                )
            )
        except RowError as e:
            errors.append({"row": i, "code": e.code})  # само содержимое строки не возвращаем

    return _finish(txs, errors, "balance" in cols)


def _make_tx(row_no: int, d: date, signed: int, counterparty: str, inn: str, purpose: str, balance: int | None, sid: str = "") -> dict:
    direction = "in" if signed > 0 else "out"
    counterparty = counterparty.strip()[:300]
    purpose = purpose.strip()[:1000]
    cat, conf, method = categorize_ex(purpose, direction, counterparty)
    cl = classify(purpose, counterparty, signed, balance)  # подкатегория, группа, документ-основание, «спасательное» для ML и ledger
    tx = {
        "row_no": row_no,
        "date": d.isoformat(),
        "amount_kop": str(signed),
        "direction": direction,
        "counterparty": counterparty,
        "inn": re.sub(r"\D", "", inn)[:12],
        "purpose": purpose,
        "balance_kop": None if balance is None else str(balance),
        "category": cat,
        "category_confidence": conf,
        "category_method": method,  # rule | rule_gleb | ml | fallback
        "subcategory": cl["subcategory"],
        "method": cl["method"],  # способ разметки подкатегории: rule | fallback (дальше ledger: tail | tenant_rule | manual)
        "group": cl["group"],
        "doc_no": cl["doc_no"],
        "doc_date": cl["doc_date"],
        "is_rescue": cl["is_rescue"],
    }
    if sid:
        tx["source_id"] = sid
    return tx


def _finish(txs: list[dict], errors: list[dict], has_balance: bool) -> dict:
    """Общий хвост разбора: модель на «хвост», спасательные поступления, режим налогообложения, статистика."""
    ml_labeled = apply_tail_model(txs)
    rescue = 0
    for t in txs:
        # собственные деньги пришли на счёт, и без них остаток дня ушёл бы в минус (признак gleb)
        t["is_rescue"] = bool(t.get("is_rescue")) or bool(
            has_balance
            and t["category"] == "transfers"
            and t["direction"] == "in"
            and t["balance_kop"] is not None
            and int(t["balance_kop"]) - int(t["amount_kop"]) < 0
        )
        rescue += t["is_rescue"]
    stats = _stats(txs, errors, has_balance)
    stats["rescue_inflows"] = rescue
    stats["rule_coverage"] = round(sum(t["category_method"] in ("rule", "rule_gleb") for t in txs) / len(txs), 4) if txs else None
    stats["ml_categorized"] = ml_labeled
    stats["tax_regime_guess"] = guess_tax_regime(txs)
    return {"transactions": txs, "errors": errors, "stats": stats}


def _stats(txs: list[dict], errors: list[dict], has_balance: bool) -> dict:
    stats = {"rows_ok": len(txs), "rows_rejected": len(errors), "has_balance": has_balance}
    if not txs:
        return stats
    ordered = sorted(txs, key=lambda t: (t["date"], t["row_no"]))
    ds = [t["date"] for t in ordered]
    d0, d1 = date.fromisoformat(ds[0]), date.fromisoformat(ds[-1])
    stats["date_from"], stats["date_to"] = ds[0], ds[-1]
    stats["days_span"] = (d1 - d0).days + 1

    # пропуски: разрыв между соседними датами операций > 5 календарных дней
    gaps, gap_days, prev = [], 0, None
    for dd in sorted(set(ds)):
        cur = date.fromisoformat(dd)
        if prev is not None and (cur - prev).days > 5:
            gaps.append([prev.isoformat(), dd])
            gap_days += (cur - prev).days - 1
        prev = cur
    stats["gap_days"] = gap_days
    stats["missing_period_count"] = len(gaps)
    stats["first_gap_dates"] = gaps[:5]

    # вероятные точные дубли (не удаляем; импортёр дедуплицирует по номеру вхождения)
    seen: set = set()
    dups = 0
    for t in ordered:
        k = (t["date"], t["amount_kop"], t["inn"], t["purpose"], t["counterparty"])
        if k in seen:
            dups += 1
        else:
            seen.add(k)
    stats["suspected_duplicates"] = dups

    # сверка остатков
    days: dict[str, list] = {}
    for t in ordered:
        if t["balance_kop"] is not None:
            days.setdefault(t["date"], []).append(t)
    checked = mism = 0
    mode = None
    opening = None
    if days:
        running = any(len({x["balance_kop"] for x in v}) > 1 for v in days.values())
        mode = "running" if running else "eod"
        if mode == "eod":
            net: dict[str, int] = {}
            for t in ordered:
                net[t["date"]] = net.get(t["date"], 0) + int(t["amount_kop"])
            pb = None
            for dd in sorted(days):
                b = int(days[dd][-1]["balance_kop"])
                if pb is None:
                    opening = b - net[dd]
                else:
                    checked += 1
                    if b - pb != net[dd]:
                        mism += 1
                pb = b
        else:
            pb = None
            for t in ordered:
                if t["balance_kop"] is None:
                    continue
                b = int(t["balance_kop"])
                if pb is None:
                    opening = b - int(t["amount_kop"])
                else:
                    checked += 1
                    if b - pb != int(t["amount_kop"]):
                        mism += 1
                pb = b
    stats["balance_mode"] = mode
    stats["balance_checked"] = checked
    stats["balance_mismatches"] = mism
    stats["opening_balance_kop"] = None if opening is None else str(opening)
    return stats


# ---------------------------------------------------------------------------
# 1CClientBankExchange (текстовый обмен «клиент-банк» 1С). Идея: проект gleb (swimshadow).
# ---------------------------------------------------------------------------
_1C_MAGIC = b"1CClientBankExchange"


def is_1c(data: bytes) -> bool:
    return data.lstrip(b"\xef\xbb\xbf \t\r\n")[: len(_1C_MAGIC)] == _1C_MAGIC


def decode_1c(data: bytes) -> str:
    """UTF-8 (с BOM или без) -> иначе cp1251; если в шапке «Кодировка=DOS» — cp866."""
    data = data[3:] if data.startswith(b"\xef\xbb\xbf") else data
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        pass
    head = data[:2048]
    # ключ «Кодировка» лежит в тех же байтах, что и весь файл: ищем его в обеих кодировках
    dos = any(
        (m := re.search(r"Кодировка\s*=\s*(\w+)", head.decode(e, errors="replace"))) and m.group(1).upper() == "DOS"
        for e in ("cp866", "cp1251")
    )
    enc = "cp866" if dos else "cp1251"
    try:
        return data.decode(enc)
    except UnicodeDecodeError:
        raise ValueError("bad_encoding")


def _acc(v: str | None) -> str:
    return re.sub(r"\s+", "", v or "")


def _cp_from_doc(doc: dict, outgoing: bool) -> tuple[str, str]:
    """(имя контрагента, ИНН). Контрагент — получатель при списании, плательщик при поступлении."""
    side = "Получатель" if outgoing else "Плательщик"
    inn = re.sub(r"\D", "", doc.get(side + "ИНН", ""))
    name = (doc.get(side + "1") or doc.get(side) or "").strip()
    m = re.match(r"^ИНН\s*(\d{10,12})\s*(.*)$", name, re.I)  # «ИНН 7701234567 ООО Ромашка» в поле без «1»
    if m:
        inn = inn or m.group(1)
        name = m.group(2).strip()
    return name, (inn if len(inn) in (10, 12) else "")


def parse_1c(data: bytes) -> dict:
    if not is_1c(data):
        raise ValueError("not_1c_format")
    text = decode_1c(data)
    header: dict[str, str] = {}
    accounts: list[dict] = []   # секции СекцияРасчСчет
    docs: list[tuple[int, dict | None]] = []   # (номер документа по порядку, поля | None если оборван)
    cur: dict | None = None
    cur_acc: dict | None = None
    seen_end = False
    for line in text.splitlines():
        line = line.strip()
        if not line:
            continue
        if line.startswith("СекцияДокумент"):
            if cur is not None:
                docs.append((len(docs) + 1, None))  # предыдущий не закрыт -> оборван
            cur = {"_type": line.partition("=")[2].strip()}
            if len(docs) > MAX_ROWS:
                raise ValueError("too_many_rows")
            continue
        if line == "КонецДокумента":
            if cur is not None:
                docs.append((len(docs) + 1, cur))
                cur = None
            continue
        if line == "СекцияРасчСчет":
            cur_acc = {}
            continue
        if line == "КонецРасчСчет":
            if cur_acc is not None:
                accounts.append(cur_acc)
            cur_acc = None
            continue
        if line == "КонецФайла":
            seen_end = True
            break
        k, eq, v = line.partition("=")
        if not eq:
            continue
        k, v = k.strip(), v.strip()
        if cur is not None:
            cur.setdefault(k, v)
        elif cur_acc is not None:
            cur_acc.setdefault(k, v)
        else:
            header.setdefault(k, v)
    if cur is not None:
        docs.append((len(docs) + 1, None))  # файл оборван внутри документа

    main = accounts[0] if accounts else {}
    account = _acc(main.get("РасчСчет") or header.get("РасчСчет"))
    if not account:
        raise ValueError("no_account")
    if not docs:
        raise ValueError("empty_file")

    txs: list[dict] = []
    errors: list[dict] = []
    seen_docs: set = set()
    for n, doc in docs:
        if doc is None:
            errors.append({"row": n, "code": "truncated_document"})
            continue
        try:
            payer = {_acc(doc.get("ПлательщикСчет")), _acc(doc.get("ПлательщикРасчСчет"))} - {""}
            recv = {_acc(doc.get("ПолучательСчет")), _acc(doc.get("ПолучательРасчСчет"))} - {""}
            is_out, is_in = account in payer, account in recv
            if is_out and is_in:
                raise RowError("ambiguous_direction")
            if not (is_out or is_in):
                raise RowError("foreign_account")  # документ не по нашему счёту: направление не угадываем
            raw_date = (doc.get("ДатаСписано") if is_out else doc.get("ДатаПоступило")) or doc.get("Дата")
            d = parse_date(raw_date)
            amount = parse_money(doc.get("Сумма"))
            if amount < 0:
                raise RowError("bad_amount")
            if amount == 0:
                raise RowError("zero_amount")
            signed = -amount if is_out else amount
            name, inn = _cp_from_doc(doc, is_out)
            purpose = doc.get("НазначениеПлатежа", "")
            num = doc.get("Номер", "").strip()
            if num:
                key = (doc["_type"], num, d, signed, tuple(sorted(payer)), tuple(sorted(recv)), purpose)
                if key in seen_docs:
                    raise RowError("duplicate_document")
                seen_docs.add(key)
            txs.append(_make_tx(n, d, signed, name, inn, purpose, None, f"{num}@{d.isoformat()}"[:100] if num else ""))
        except RowError as e:
            errors.append({"row": n, "code": e.code})

    # остатки: начало дня из НачальныйОстаток секции нашего счёта + оборот; сверка с КонечныйОстаток
    opening = closing = None
    for key in ("НачальныйОстаток", "КонечныйОстаток"):
        raw = main.get(key)
        if raw:
            try:
                val = parse_money(raw)
            except RowError:
                continue
            if key == "НачальныйОстаток":
                opening = val
            else:
                closing = val
    has_balance = opening is not None and bool(txs)
    if has_balance:
        txs.sort(key=lambda t: (t["date"], t["row_no"]))
        running = opening
        i = 0
        while i < len(txs):
            j = i
            while j < len(txs) and txs[j]["date"] == txs[i]["date"]:
                running += int(txs[j]["amount_kop"])
                j += 1
            for t in txs[i:j]:
                t["balance_kop"] = str(running)
            i = j
    res = _finish(txs, errors, has_balance)
    st = res["stats"]
    st["source_format"] = "1c"
    st["file_complete"] = seen_end
    st["closing_balance_kop"] = None if closing is None else str(closing)
    if has_balance and closing is not None:
        ok = running == closing
        st["closing_balance_ok"] = ok
        st["balance_checked"] += 1
        st["balance_mismatches"] += 0 if ok else 1
    else:
        st["closing_balance_ok"] = None
    return res
